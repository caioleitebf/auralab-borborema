"""Verifica se ha imagens incorporadas no API.xlsx."""
import openpyxl
import zipfile
from pathlib import Path


def main():
    f = Path(r"C:\Users\caio.ferreira\Downloads\API (1).xlsx")

    # 1) openpyxl
    wb = openpyxl.load_workbook(f)
    ws = wb.active
    print(f"Sheet ativa: {ws.title}")
    print(f"  Imagens (openpyxl): {len(ws._images)}")
    print(f"  Charts: {len(ws._charts)}")
    print(f"  Tables: {len(ws.tables) if hasattr(ws, 'tables') else 0}")
    for i, img in enumerate(ws._images[:30]):
        try:
            anchor = img.anchor
            r = anchor.from_.row + 1
            c = anchor.from_.col + 1
            print(f"  Img {i}: linha {r}, col {c} ({openpyxl.utils.get_column_letter(c)})")
        except Exception as e:
            print(f"  Img {i}: erro lendo anchor: {e}")

    # 2) Inspecao raw via zip (xlsx e um zip)
    print()
    print("=== Conteudo do XLSX (como zip) ===")
    with zipfile.ZipFile(f) as z:
        for name in z.namelist():
            info = z.getinfo(name)
            print(f"  {name:60s} ({info.file_size:>9} bytes)")

    # 3) Extrai imagens encontradas para uma pasta para podermos ver
    out_dir = Path(__file__).parent / "_api_xlsx_imgs"
    out_dir.mkdir(exist_ok=True)
    print()
    print(f"=== Extraindo midia/imagens para {out_dir} ===")
    with zipfile.ZipFile(f) as z:
        for name in z.namelist():
            if "media/" in name.lower() or "image" in name.lower():
                dest = out_dir / Path(name).name
                with z.open(name) as src, open(dest, "wb") as dst:
                    dst.write(src.read())
                print(f"  -> {dest.name} ({dest.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
